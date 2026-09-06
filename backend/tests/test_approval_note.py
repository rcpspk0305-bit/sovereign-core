"""Unit and integration tests for Approval-Note Artifact Generator."""

from pathlib import Path
import pytest
from docx import Document

from app.core.agents.inspection_agent import InspectionAnalysisAgent
from app.core.tools.approval_note import ApprovalNoteGeneratorTool
from app.core.tools.registry import ControlledToolRegistry
from tests.conftest import MockLLMClient


@pytest.fixture
def temp_artifacts_dir(tmp_path):
    """Temporary directory for artifacts."""
    artifacts = tmp_path / "artifacts"
    artifacts.mkdir()
    return artifacts


@pytest.fixture
def approval_tool(temp_artifacts_dir):
    """ApprovalNoteGeneratorTool instance with temp dir."""
    return ApprovalNoteGeneratorTool(artifacts_dir=temp_artifacts_dir)


@pytest.mark.asyncio
async def test_generate_approval_note_with_verified_claims(approval_tool):
    """Test generating a verified approval note produces valid DOCX and human approval block."""
    result = await approval_tool.execute(
        title="Mission Approval Note: Apollo99 Thermal Variance Assessment",
        decision="CONDITIONAL_APPROVAL",
        summary="Inspection confirms nominal telemetry with a 7 degree Celsius temperature delta.",
        findings=[
            {
                "statement": "Apollo99 measured temperature is 82 C at t=140s.",
                "citation": "apollo99_manual.pdf (Page 4)",
            },
            {
                "statement": "Baseline operating limit is 75 C under nominal load.",
                "citation": "apollo99_manual.pdf (Page 2)",
            },
        ],
        citations=[
            "apollo99_manual.pdf (Page 4): Sensor reads 82 C at t=140s.",
            "apollo99_manual.pdf (Page 2): Baseline limit is 75 C.",
        ],
        retrieved_evidence="Apollo99 sensor readings indicate 82 C. The nominal baseline is 75 C.",
        risk_assessment="Slight temperature variance within 10% thermal tolerance.",
        human_approval_role="Chief Mission Systems Engineer",
    )

    assert result.success is True
    out = result.output
    assert out["title"] == "Mission Approval Note: Apollo99 Thermal Variance Assessment"
    assert out["decision"] == "CONDITIONAL_APPROVAL"
    assert out["total_claims"] == 2
    assert out["verified_claims_count"] == 2
    assert out["unsupported_claims_count"] == 0
    assert "FULLY GROUNDED" in out["validation_status"]

    # Verify DOCX was created locally
    docx_path = Path(out["docx_file_path"])
    assert docx_path.exists()
    assert out["docx_file_size_bytes"] > 0
    assert out["checksum_sha256"] is not None

    # Inspect DOCX contents
    doc = Document(str(docx_path))
    table_text = " ".join([c.text for t in doc.tables for r in t.rows for c in r.cells])
    doc_text = " ".join([p.text for p in doc.paragraphs]) + " " + table_text
    assert "Executive Summary" in doc_text
    assert "Explicit Human Approval & Governance Sign-Off" in doc_text
    assert "Chief Mission Systems Engineer" in doc_text
    assert "VERIFIED" in table_text


@pytest.mark.asyncio
async def test_unsupported_claim_validation_and_flagging(approval_tool):
    """Test that unsupported claims are flagged and not silently presented as verified facts."""
    result = await approval_tool.execute(
        title="Incident Report: High Risk Unauthorized Assertion",
        decision="REJECTED",
        summary="Evaluation of claims with ungrounded metrics.",
        findings=[
            {
                "statement": "Core temperature reached 999 C resulting in hull breach.",
                # No citation and 999 C is NOT in the retrieved evidence
            },
            {
                "statement": "Standard baseline is 75 C.",
                "citation": "manual.pdf (Page 1)",
            },
        ],
        citations=["manual.pdf (Page 1): Baseline is 75 C."],
        retrieved_evidence="Baseline operating temperature is 75 C.",
        human_approval_role="Safety Auditor",
    )

    assert result.success is True
    out = result.output
    assert out["total_claims"] == 2
    assert out["verified_claims_count"] == 1
    assert out["unsupported_claims_count"] == 1
    assert "FLAGGED" in out["validation_status"]
    assert "Core temperature reached 999 C resulting in hull breach." in out["unsupported_claims"]

    # Verify warning is in Markdown
    assert "UNSUPPORTED CLAIMS DETECTED" in out["document_content"]
    assert "[UNVERIFIED CLAIM]" in out["document_content"]

    # Verify warning is in DOCX
    docx_path = Path(out["docx_file_path"])
    assert docx_path.exists()
    doc = Document(str(docx_path))
    doc_text = " ".join([p.text for p in doc.paragraphs])
    assert "UNSUPPORTED CLAIMS IDENTIFIED" in doc_text
    assert "UNVERIFIED" in doc_text


def test_controlled_tool_registry_integration():
    """Verify tool registration in ControlledToolRegistry."""
    registry = ControlledToolRegistry()
    tool = ApprovalNoteGeneratorTool()
    registry.register(tool)

    registered = registry.get("approval_note_generator")
    assert registered is not None
    assert registered.name == "approval_note_generator"

    defn = registered.get_definition()
    assert defn.name == "approval_note_generator"
    assert "findings" in defn.parameters["properties"]
    assert "human_approval_role" in defn.parameters["properties"]


@pytest.mark.asyncio
async def test_agent_invokes_approval_note_generator():
    """Verify InspectionAnalysisAgent recognizes and executes approval_note_generator."""
    registry = ControlledToolRegistry()
    tool = ApprovalNoteGeneratorTool()
    registry.register(tool)

    # Mock model calling approval_note_generator
    call_json = (
        '```json\n'
        '{\n'
        '  "thought": "Generating formal approval note with citations.",\n'
        '  "tool": "approval_note_generator",\n'
        '  "arguments": {\n'
        '    "title": "Agent Mission Approval Note",\n'
        '    "summary": "Agent completed inspection successfully.",\n'
        '    "findings": [\n'
        '      {"statement": "System operating at 75 C.", "citation": "spec.pdf (Page 1)"}\n'
        '    ],\n'
        '    "citations": ["spec.pdf (Page 1)"]\n'
        '  }\n'
        '}\n'
        '```'
    )
    llm = MockLLMClient(response_text=call_json)
    agent = InspectionAnalysisAgent(llm_client=llm, tool_registry=registry, max_allowed_steps=2)

    emitted_events = []
    async def callback(ev):
        emitted_events.append(ev)

    res = await agent.run(
        prompt="Generate an approval note for system operating temperature.",
        max_steps=1,
        event_callback=callback,
    )

    assert res.success is True
    # Verify tool was called
    tool_steps = [s for s in res.steps if s.tool_name == "approval_note_generator"]
    assert len(tool_steps) == 1
    assert tool_steps[0].tool_result.success is True
    assert tool_steps[0].tool_result.output["artifact_type"] == "approval_note"

    # Verify artifact_generated was emitted
    art_events = [e for e in emitted_events if e["type"] == "artifact_generated"]
    assert len(art_events) == 1
    assert art_events[0]["artifact"]["artifact_type"] == "approval_note"
    assert art_events[0]["artifact"]["docx_file_path"] is not None
